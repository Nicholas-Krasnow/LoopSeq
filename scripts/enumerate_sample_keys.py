#!/usr/bin/env python3
"""
Read sample_keys.xlsx and print one TSV line per sample for batch submission.

Fields: sample_name, input_fasta, reference_fasta, gene_length, gene_start

Sheet "1-48": table coordinates A1:H12 map to Excel B2:M9 via
  excel_row = 2 + tc, excel_col = 2 + tr
  (tr = table row 0..11 for rows 1..12, tc = table col 0..7 for A..H).
Sample order: row-major in table coords (A1, B1, …, H1, A2, …) → sample_num = tr*8 + tc + 1.

Sheet "initial_set": one row per sample using Folder, gene_start, reference_fasta columns.
"""

from __future__ import annotations

import argparse
import re
import sys
from pathlib import Path

try:
    from openpyxl import load_workbook
except ImportError:
    print("Error: openpyxl is required (pip install openpyxl)", file=sys.stderr)
    sys.exit(1)

GENE_LENGTH_GRID = 1233


def read_first_fasta_sequence(path: Path) -> str:
    parts: list[str] = []
    with open(path, encoding="utf-8", errors="replace") as f:
        for line in f:
            line = line.strip()
            if not line or line.startswith(">"):
                if parts and line.startswith(">"):
                    break
                continue
            parts.append(line)
    return "".join(parts).replace(" ", "").upper()


def cell_str(val) -> str:
    if val is None:
        return ""
    if isinstance(val, float) and val == int(val):
        return str(int(val))
    return str(val).strip()


def pick_sp23_sp24(text: str) -> str | None:
    s = text.lower()
    if "sp24" in s:
        return "sp24"
    if "sp23" in s:
        return "sp23"
    return None


def find_trimmed_fa_under_broad(samples_root: Path, sample_num: int) -> Path | None:
    n = sample_num
    suffix = f"_BROAD_{n}"
    for out_dir in sorted(samples_root.glob("*_output")):
        if not out_dir.is_dir():
            continue
        for broad in sorted(out_dir.iterdir()):
            if not broad.is_dir() or not broad.name.endswith(suffix):
                continue
            hits = sorted(
                p for p in broad.rglob("*") if p.is_file() and p.name.endswith("trimmed.fa")
            )
            if hits:
                return hits[0]
    return None


def input_fasta_for_sample_number(base: Path, sample_num: int) -> Path | None:
    if 1 <= sample_num <= 24:
        root = base / "samples_1-24"
    elif 25 <= sample_num <= 48:
        root = base / "samples_25-48"
    else:
        return None
    if not root.is_dir():
        return None
    return find_trimmed_fa_under_broad(root, sample_num)


def emit_line(sample_name: str, inp: Path, ref: Path, gene_start: str, gene_length: int) -> None:
    # TSV; gene_start last so simple parsers can avoid embedded tabs in earlier fields
    row = [
        sample_name,
        str(inp.resolve()),
        str(ref.resolve()),
        str(gene_length),
        gene_start.replace("\t", " ").replace("\n", ""),
    ]
    print("\t".join(row))


def iter_sheet_1_48(base: Path, ws) -> None:
    ref_dir = base / "reference_fastas"
    sp23_ref = ref_dir / "sp23_dna.fa"
    sp24_ref = ref_dir / "sp24_dna.fa"
    sp23_start = read_first_fasta_sequence(ref_dir / "sp23_gene_start.fa")
    sp24_start = read_first_fasta_sequence(ref_dir / "sp24_gene_start.fa")

    col_letters = "ABCDEFGH"
    for tr in range(12):
        for tc in range(8):
            excel_row = 2 + tc
            excel_col = 2 + tr
            cell = ws.cell(row=excel_row, column=excel_col)
            text = cell_str(cell.value)
            if not text:
                continue
            sample_num = tr * 8 + tc + 1
            coord = f"{col_letters[tc]}{tr + 1}"
            sample_name = f"1-48_{coord}_n{sample_num}"

            which = pick_sp23_sp24(text)
            if which is None:
                print(
                    f"skip 1-48 {sample_name}: no sp23/sp24 in cell text",
                    file=sys.stderr,
                )
                continue

            if which == "sp23":
                ref_fa, gstart = sp23_ref, sp23_start
            else:
                ref_fa, gstart = sp24_ref, sp24_start

            if not ref_fa.is_file():
                print(f"skip {sample_name}: missing {ref_fa}", file=sys.stderr)
                continue

            inp = input_fasta_for_sample_number(base, sample_num)
            if inp is None or not inp.is_file():
                print(
                    f"skip {sample_name}: no trimmed.fa for BROAD_{sample_num}",
                    file=sys.stderr,
                )
                continue

            emit_line(sample_name, inp, ref_fa, gstart, GENE_LENGTH_GRID)


def normalize_header(name: str) -> str:
    return re.sub(r"\s+", "_", name.strip().lower())


def iter_initial_set(base: Path, ws) -> None:
    ref_dir = base / "reference_fastas"
    initial_root = base / "initial_set"

    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return
    headers = [normalize_header(cell_str(c)) for c in rows[0]]
    try:
        ix_folder = headers.index("folder")
        ix_gene = headers.index("gene_start")
        ix_ref = headers.index("reference_fasta")
    except ValueError as e:
        print(f"initial_set: missing column {e}", file=sys.stderr)
        return

    for i, row in enumerate(rows[1:], start=2):
        if not row:
            continue
        def get(j):
            return cell_str(row[j]) if j < len(row) else ""

        folder = get(ix_folder)
        gene_start = get(ix_gene)
        ref_name = get(ix_ref)
        if not folder or not gene_start or not ref_name:
            continue

        sample_name = f"initial_set_r{i}_{re.sub(r'[^A-Za-z0-9._-]+', '_', folder)[:80]}"

        ref_path = ref_dir / ref_name
        if not ref_path.is_file():
            print(f"skip {sample_name}: reference not found {ref_path}", file=sys.stderr)
            continue

        inp: Path | None = None
        if initial_root.is_dir():
            for out_dir in sorted(initial_root.glob("*_output")):
                if not out_dir.is_dir():
                    continue
                for sub in sorted(out_dir.iterdir()):
                    if not sub.is_dir() or not sub.name.endswith(folder):
                        continue
                    hits = sorted(
                        p
                        for p in sub.rglob("*")
                        if p.is_file() and p.name.endswith("_trimmed.fa")
                    )
                    if hits:
                        inp = hits[0]
                        break
                if inp:
                    break

        if inp is None or not inp.is_file():
            print(
                f"skip {sample_name}: no *_trimmed.fa under initial_set for Folder={folder!r}",
                file=sys.stderr,
            )
            continue

        emit_line(sample_name, inp, ref_path, gene_start, 1233)


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument(
        "base_dir",
        type=Path,
        help="Project root (contains sample_keys.xlsx, samples_*, reference_fastas, …)",
    )
    ap.add_argument(
        "--xlsx",
        type=Path,
        default=None,
        help="Path to workbook (default: <base_dir>/sample_keys.xlsx)",
    )
    args = ap.parse_args()
    base = args.base_dir.resolve()
    xlsx = args.xlsx or (base / "sample_keys.xlsx")
    if not xlsx.is_file():
        print(f"Error: workbook not found: {xlsx}", file=sys.stderr)
        sys.exit(1)

    wb = load_workbook(xlsx, read_only=True, data_only=True)

    if "1-48" in wb.sheetnames:
        iter_sheet_1_48(base, wb["1-48"])
    else:
        print('Warning: sheet "1-48" not found', file=sys.stderr)

    if "initial_set" in wb.sheetnames:
        iter_initial_set(base, wb["initial_set"])
    else:
        print('Warning: sheet "initial_set" not found', file=sys.stderr)

    wb.close()


if __name__ == "__main__":
    main()
